"""Create the four-page report from detailed-study evidence, never smoke scores.

Run with the document runtime Python (python-docx + Pillow). Render the resulting
DOCX and inspect every page before submitting. Does not train or evaluate models.
"""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
STUDY = ROOT / "artifacts_retrain_v3"
OUT = ROOT / "submission/COS30082_Report.docx"
QA = ROOT / "tmp/report_qa"

def read_json(path): return json.loads(path.read_text(encoding="utf-8"))
def read_csv(path):
    with path.open(encoding="utf-8", newline="") as f: return list(csv.DictReader(f))
def pct(x): return f"{100 * float(x):.2f}"

def build():
    QA.mkdir(parents=True, exist_ok=True)
    OUT.parent.mkdir(exist_ok=True)
    paths = sorted(STUDY.glob("*/*/result.json"))
    rows = [read_json(p) for p in paths]
    assert len(rows) == 26 and len({(r['architecture'],r['experiment']) for r in rows}) == 26
    by = {(r["architecture"],r["experiment"]):r for r in rows}
    names = read_json(STUDY / "split.json")["class_names"]
    test = read_csv(STUDY / "test_summary.csv")
    lock = read_json(STUDY / "test_lock.json")
    assert set(lock["selected"]) == {r["run"] for r in test}
    winner = max(rows, key=lambda r:r["val_top1"])
    assert winner["architecture"] == "efficientnet_v2_s" and winner["experiment"] == "03_horizontal_flip"
    win_test = next(r for r in test if r["run"] == "efficientnet_v2_s/03_horizontal_flip")
    predictions = read_csv(STUDY / win_test["run"] / "test/predictions.csv")
    assert len(predictions) == 1204
    assert sum(r["truth"] == r["prediction"] for r in predictions) == 920
    assert abs(float(win_test["top1"]) - 920/1204) < 1e-12
    assert all(int(r["overlap_count"]) == 4 and int(r["clean_samples"]) == 1200 for r in test)
    perclass = read_csv(STUDY / win_test["run"] / "test/per_class.csv")
    assert len(perclass) == 200
    perfect = sum(float(r["accuracy"]) == 1 for r in perclass)
    zero = sum(float(r["accuracy"]) == 0 for r in perclass)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
    sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
    sec.footer_distance = Inches(.28)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0,0,0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].font.size = Pt(11)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].paragraph_format.space_before = Pt(4)
    doc.styles["Caption"].paragraph_format.space_after = Pt(6)
    # The bundled Word template contains a blue Title border; remove all
    # paragraph borders so the report uses plain black headings consistently.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    def p(text, style=None):
        preceding = doc.element.body[-2] if len(doc.element.body) > 1 else None
        para = doc.add_paragraph(text, style)
        if preceding is not None and preceding.tag == qn("w:tbl"):
            para.paragraph_format.space_before = Pt(6)
        return para
    def h(text): doc.add_heading(text,2)
    def caption(text): p(text,"Caption")
    def table(headers, data, widths):
        t = doc.add_table(rows=1,cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for c,w in zip(t.columns,widths): c.width = Inches(w)
        for c,s in zip(t.rows[0].cells,headers): c.text = s
        for row in data:
            for c,s in zip(t.add_row().cells,row): c.text = str(s)
        borders = OxmlElement("w:tblBorders")
        for edge in ["top","left","bottom","right","insideH","insideV"]:
            el=OxmlElement(f"w:{edge}")
            for k,v in {"val":"single","sz":"4","color":"D9D9D9"}.items(): el.set(qn(f"w:{k}"),v)
            borders.append(el)
        t._tbl.tblPr.append(borders)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for ri,row in enumerate(t.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for ci,c in enumerate(row.cells):
                c.width = Inches(widths[ci])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = c._tc.get_or_add_tcPr()
                fill=OxmlElement("w:shd")
                fill.set(qn("w:fill"),"243746" if ri==0 else ("F2F5F7" if ri%2 else "FFFFFF"))
                pr.append(fill)
                margins=OxmlElement("w:tcMar")
                for edge in ["top","left","bottom","right"]:
                    e=OxmlElement(f"w:{edge}");e.set(qn("w:w"),"70");e.set(qn("w:type"),"dxa");margins.append(e)
                pr.append(margins)
                for para in c.paragraphs:
                    para.paragraph_format.space_before=Pt(0)
                    para.paragraph_format.space_after=Pt(0)
                    para.paragraph_format.line_spacing=1.0
                    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.name="Calibri";run.font.size=Pt(9)
                        run.font.bold=ri==0
                        run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
        return t
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Fa Zhen NG | 104400141 | ").font.size=Pt(8)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)

    # Page 1: methodology.
    p("Bird Species Classification with Transfer Learning","Title")
    p("COS30082 Assignment 1 | Fa Zhen NG | Student ID 104400141")
    doc.add_heading("Methodology",1)
    p("I compared MobileNetV1 and EfficientNetV2-S for 200-class bird identification using 13 cumulative training recipes per architecture. The strongest validation recipe was EfficientNetV2-S with backbone fine-tuning, random resized crops and horizontal flips. It achieved 76.41% Top-1 accuracy and 75.90% average accuracy per class on the official Test set. Additional aggressive augmentation did not improve the selected model.")
    h("Data preparation and evaluation")
    p("The assignment provides 4,829 Train images and 1,204 Test images with labels 0-199. I verified files and hashed their contents. Three duplicate-content pairs in Train have conflicting labels; all six rows remained in fit with their original labels. A seed-42 stratified split of the remaining unique images produced 3,863 fit and 966 validation images, retaining all 200 species in both partitions without identical content crossing the boundary.")
    p("Validation selected checkpoints and recipes. Test was evaluated only after the 26 runs, with selections saved in a lock file. Four Test images duplicate Train content; I report the official 1,204-image results and a supplementary 1,200-image subset excluding those overlaps. Top-1 is the proportion of correctly classified images. Average accuracy per class is the mean recall across the 200 species, so each species receives equal weight.")
    h("CNN architectures and input processing")
    p("MobileNetV1 width 1.0 uses depthwise spatial convolutions followed by 1 x 1 pointwise convolutions, reducing computation [1]. The timm model mobilenetv1_100.ra4_e3600_r224_in1k supplies ImageNet weights. EfficientNetV2-S combines Fused-MBConv and MBConv blocks and uses torchvision ImageNet weights [2]. Both feature extractors feed global average pooling and a new 200-output linear classifier. Native stochastic depth is disabled; classifier dropout is introduced only at its designated stage.")
    p("Images are converted to RGB and resized to 224 x 224. MobileNet uses channel mean/std 0.5; EfficientNet uses ImageNet mean (0.485, 0.456, 0.406) and std (0.229, 0.224, 0.225). Validation, Test and deployment share deterministic preprocessing. Augmentation is resampled only when fit images are loaded.")
    h("Training scheme and experimental control")
    p("Training uses AdamW, head LR 3e-4, backbone LR 3e-5, cosine decay after a two-epoch classifier warm-up, a 15-epoch maximum, and validation Top-1 early stopping with patience 5. Batch 8 with accumulation 4 gives a nominal effective batch of 32; AMP runs on an RTX 3050 Laptop GPU with 4 GB VRAM. BatchNorm running statistics stay fixed. Baselines train only the new head; later recipes unfreeze backbone weights after warm-up. The loss is cross-entropy, with smoothing added at stage 10. Displayed losses remain unsmoothed cross-entropy.")
    p("Every recipe restarts from the same architecture-specific ImageNet weights and seed; it does not continue the previous bird checkpoint. Split, resolution, epoch budget and selection policy are shared. Later rows retain earlier methods, so adjacent differences measure conditional additions rather than independent effects. Pretraining histories differ between architectures, and one seed does not establish statistical significance.")

    # Page 2: full paired stage results.
    doc.add_page_break()
    doc.add_heading("Result and Discussion",1)
    h("Cumulative validation experiments")
    labels=["Frozen backbone baseline","Backbone fine-tuning","Random resized crop","Horizontal flip","Vertical flip","Random rotation","ColorJitter","Inversion and colour layers","Classifier dropout","Weight decay","Label smoothing","Random erasing","Grayscale"]
    stages=sorted({r['experiment'] for r in rows})
    data=[]
    for i,(stage,label) in enumerate(zip(stages,labels)):
        m=by['mobilenet_v1',stage];e=by['efficientnet_v2_s',stage]
        dm=0 if not i else 100*(m['val_top1']-by['mobilenet_v1',stages[i-1]]['val_top1'])
        de=0 if not i else 100*(e['val_top1']-by['efficientnet_v2_s',stages[i-1]]['val_top1'])
        data.append([f"{i:02d} {label}",pct(m['val_top1']),pct(m['val_average_per_class']),pct(e['val_top1']),pct(e['val_average_per_class']),f"{dm:+.2f}",f"{de:+.2f}"])
    caption("Table 1. Best validation checkpoint of each recipe. M = MobileNetV1; E = EfficientNetV2-S; Avg = average accuracy per class. Delta is change from the preceding recipe in percentage points (pp). All accuracy columns are percentages.")
    table(["Cumulative recipe","M Top-1","M Avg","E Top-1","E Avg","Delta M\npp","Delta E\npp"],data,[2.2,.8,.8,.8,.8,.8,.8])
    p("The crop stage uses area 40-100% and aspect ratio 0.75-1.33; horizontal and vertical flips use probabilities 0.5 and 0.2. Rotation is uniform from -45 to +45 degrees. ColorJitter uses brightness/contrast 0.4, saturation 0.5 and hue 0.1. Colour effects trigger with probability 0.25, selecting inversion or red/green/blue/yellow overlays; tint opacity is 0.20-0.45. Later settings add dropout 0.3, weight decay 0.001, smoothing 0.1, erasing probability 0.25 with area 2-15%, and grayscale probability 0.2.")
    p("The frozen MobileNet baseline exceeded EfficientNet by 9.32 pp. Fine-tuning changed that ranking: EfficientNet gained 19.77 pp, compared with 1.14 pp for MobileNet. Adding crop and horizontal flip brought EfficientNet to its best 76.09% validation Top-1, 25.05 pp above its baseline. MobileNet's best was the fine-tuning-only recipe at 61.49%; cropping did not improve it under these settings.")
    p("Stronger transformations were harmful. MobileNet fell from 55.69% after rotation to 48.76% with ColorJitter and 45.13% with colour effects. EfficientNet also declined, but less sharply. A plausible explanation is that altered feather colours or unnatural orientations remove useful species cues; these experiments do not establish that mechanism. Regularization stages partly recovered EfficientNet, but none surpassed its earlier crop/flip recipe. The final grayscale recipes reached only 41.20% and 72.98%. More augmentation therefore did not imply better accuracy.")

    # Page 3: final Test and overfitting with quantitative evidence.
    doc.add_page_break()
    h("Final Test performance and model selection")
    caption("Table 2. Official Test (N = 1,204) and clean Test (N = 1,200). Recipes were chosen using validation, before Test evaluation. All values are percentages.")
    display_names={"mobilenet_v1/00_baseline":"MobileNetV1 baseline","mobilenet_v1/01_fine_tune":"MobileNetV1 fine-tuning","efficientnet_v2_s/00_baseline":"EfficientNetV2-S baseline","efficientnet_v2_s/03_horizontal_flip":"EfficientNetV2-S crop and flip"}
    table(["Selected model","Test\nTop-1","Test\nAvg/class","Clean\nTop-1","Clean\nAvg/class"],[[display_names[r['run']],pct(r['top1']),pct(r['average_per_class']),pct(r['clean_top1']),pct(r['clean_average_per_class'])] for r in test],[2.4,1.15,1.15,1.15,1.15])
    p("EfficientNetV2-S with fine-tuning, crop and horizontal flip was selected for deployment by validation performance. It correctly classified 920 of 1,204 Test images (76.41%), improving its frozen baseline by 25.58 pp. It exceeded the selected MobileNet model by 17.36 pp. Average per-class accuracy was 75.90%, close to overall accuracy, although this does not imply uniform performance for every species.")
    p("Removing the four known overlaps increased the winner's Top-1 to 76.67% and per-class average to 76.17%. The small positive change means these four images did not inflate this model's measured score; excluding them still provides a clearer generalization estimate. This audit detects exact duplicates only, not near-duplicates or overlap with the original ImageNet pretraining data.")
    h("Overfitting and computational cost")
    data=[]
    for r in test:
        arch,stage=r['run'].split('/');res=by[arch,stage]
        hist=read_csv(STUDY/r['run']/"history.csv")
        row=next(x for x in hist if int(x['epoch'])==res['best_epoch'])
        data.append([display_names[r['run']],pct(row['fit_eval_top1']),pct(row['val_top1']),f"{100*float(row['gap']):.2f}",f"{float(row['train_loss']):.3f}",f"{float(row['val_loss']):.3f}",str(res['best_epoch'])])
    caption("Table 3. Metrics at the validation-selected epoch. Train Eval and Val use deterministic inputs; Gap = Train Eval minus Val. Loss columns are ordinary cross-entropy; Train loss is logged during optimization.")
    table(["Selected model","Train\nEval %","Val %","Gap\npp","Train\nloss","Val\nloss","Best\nepoch"],data,[2.2,.8,.8,.8,.8,.8,.8])
    p("Both selected models show substantial fit/validation gaps. Fine-tuned MobileNet reached 99.92% fit accuracy but only 61.49% validation, a 38.43 pp gap. EfficientNet's selected crop/flip recipe reached 99.64% fit and 76.09% validation, a 23.55 pp gap. Its training loss was 0.033 versus validation loss 1.207. These results indicate incomplete generalization despite augmentation and early checkpoint selection.")
    p("Colour augmentation reduced MobileNet's fit accuracy as well as validation accuracy, while its gap narrowed from 25.69 pp after rotation to 24.09 pp with ColorJitter. Its declining validation score cannot therefore be described solely as worsening overfitting; learning under the stronger transformations also became harder. Future validation-only studies could test weaker colour changes and more targeted crops, rather than assuming additional methods are beneficial.")
    mtime=by['mobilenet_v1','01_fine_tune']['training_seconds']/60
    etime=winner['training_seconds']/60
    p(f"The selected MobileNet and EfficientNet runs took {mtime:.1f} and {etime:.1f} minutes respectively on this laptop. Across all recipes, recorded training time totalled {sum(r['training_seconds'] for r in rows)/3600:.2f} hours, including epoch evaluation/checkpoint work but excluding initial model setup. The larger EfficientNet recipe cost more per run, but its substantial accuracy advantage justified selecting it for this application. These are training times, not inference-latency benchmarks.")

    # Page 4: observed errors, deployment and references within the second section.
    doc.add_page_break()
    h("Class errors and qualitative examples")
    p(f"The selected EfficientNet made 284 Test errors. Of 200 species, {perfect} had perfect recall and {zero} had zero recall; each class contained only 4-8 Test images, so these class scores are noisy. Repeated confusions included Boat-tailed Grackle to Brewer Blackbird (3 images), House Wren to Great Crested Flycatcher (3), and Acadian Flycatcher to Least Flycatcher (2). These are observed error counts; visual similarity is a possible interpretation rather than a measured cause.")
    grackle=names.index("Boat tailed Grackle")
    correct=next(r for r in predictions if int(r['truth'])==grackle and r['truth']==r['prediction'])
    wrong=next(r for r in predictions if int(r['truth'])==grackle and r['truth']!=r['prediction'])
    para=p("")
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i,r in enumerate([correct,wrong]):
        image=ROOT/"Test"/r['filename']
        with Image.open(image) as im: w,hgt=im.size
        factor=min(3.1/w,1.65/hgt)
        para.add_run().add_picture(str(image),width=Inches(w*factor),height=Inches(hgt*factor))
        if i==0: para.add_run("     ")
    caption(f"Figure 1. Two actual Test predictions from the selected model. Left: {correct['filename']}, Boat-tailed Grackle, correctly classified. Right: {wrong['filename']}, also Boat-tailed Grackle, predicted as {names[int(wrong['prediction'])]}. Examples illustrate a success and a failure, not a representative accuracy estimate.")
    h("Deployment and limitations")
    p("The Streamlit application loads the validation-selected checkpoint, its 200 species labels, normalization and input size. It supports photo uploads, Top-5 probabilities, checkpoint selection and experiment tables. CPU inference allows the GPU to remain available for training. I use Streamlit with instructor approval in place of the original Hugging Face demonstration. The exported model is artifacts_retrain_v3/deployment/bird_model.pth.")
    p("The application always chooses among known species: unfamiliar birds, multiple birds, distant subjects or different photographic conditions can yield incorrect confident predictions. A high score on one upload is not deployment accuracy, and no independent upload-photo benchmark was recorded. The study also uses one random seed, a small validation set and different ImageNet pretraining histories. Results support the selected recipe on this assignment split, rather than a general claim that EfficientNet is always superior.")
    p("Run the application: .\\.venv\\Scripts\\python.exe -m streamlit run app.py. Open the local URL printed by Streamlit. Source: https://github.com/FZhen2450/COS30082_individual_assignment")
    p("YouTube presentation link: [TO BE ADDED AFTER RECORDING]")
    h("References")
    p("[1] Howard, A. G., et al. (2017). MobileNets: Efficient Convolutional Neural Networks for Mobile Vision Applications. arXiv:1704.04861. https://arxiv.org/abs/1704.04861", "Caption")
    p("[2] Tan, M., and Le, Q. V. (2021). EfficientNetV2: Smaller Models and Faster Training. ICML. arXiv:2104.00298. https://arxiv.org/abs/2104.00298", "Caption")
    p("[3] COS30082 Assignment 1: Bird Species Classification. Assignment1_v2.pdf, pp. 1-3. Experimental evidence: detailed notebook and artifacts_retrain_v3 run metadata, epoch histories, test_summary.csv and saved Test predictions.", "Caption")
    doc.core_properties.author="Fa Zhen NG"
    doc.core_properties.title="Bird Species Classification with Transfer Learning"
    doc.core_properties.subject="COS30082 Assignment 1"
    doc.save(OUT)
    sources=paths+[STUDY/"test_summary.csv",STUDY/"test_lock.json",ROOT/"Assignment1_v2.pdf"]
    (QA/"report_evidence.json").write_text(json.dumps({
        "study":"artifacts_retrain_v3", "runs":len(rows), "winner":winner,
        "test_correct":920,"test_total":1204,"known_overlaps":4,
        "sources":{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},
        "youtube_link_pending":True},indent=2),encoding="utf-8")
    print(OUT)

if __name__=="__main__": build()
